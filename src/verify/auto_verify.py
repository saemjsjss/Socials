"""
Automatic verification of downloaded documents (run by the bot after each portal sync).

Each pass:
  1. Document check — every student whose folder or portal record changed since the last
     check is re-read (PDF text layer first, OCR only for scans) and tested against the
     program guidelines.  Results go into DOCUMENT CHECK.xlsx, one sheet per program.
  2. Field check — the same students' portal values are compared against what their
     documents actually say.  Results go into FIELD CHECK.xlsx, one sheet per program.

Every document is OCR'd once only: the text is kept in data/verification/text/<PASSPORT>.json
and reused until that file itself changes.  So when an employee corrects the portal after a
flag, only the (fast) field check runs again — the scans are not read a second time.

Both reports are rebuilt in full from data/verification/results.json every pass, so a
run that only checks two new students still produces a complete report for everyone.
FIELD CHECK.xlsx also carries a "Corrections" sheet: every portal field that changed after
a check, with what it was, what it is now, and whether the change fixed the problem.

OCR is slow (minutes per student), so each pass has a budget of students whose documents
are new; whatever is left over is picked up by the next pass.

CLI (run from the BOT folder):
  python -m src.verify.auto_verify                # check what is pending (default budget)
  python -m src.verify.auto_verify --budget 0     # no limit: drain the whole queue
  python -m src.verify.auto_verify --pending      # just list what is waiting
  python -m src.verify.auto_verify --recheck      # re-test everyone against the current rules
  python -m src.verify.auto_verify --rebuild      # only rewrite the two reports
  python -m src.verify.auto_verify --recheck-all  # forget every result and start over
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import logging
import os
import time
from pathlib import Path
from typing import Any, Dict, List, Tuple

from src.config import settings

logger = logging.getLogger(__name__)

REPORT_DIR = settings.verification_dir()
STORE_PATH = REPORT_DIR / "results.json"
TEXT_DIR = REPORT_DIR / "text"          # OCR output, kept so nothing is read twice
LOCK_PATH = REPORT_DIR / "auto_verify.lock"
LOCK_STALE_SECONDS = 4 * 3600
DOC_REPORT = REPORT_DIR / "DOCUMENT CHECK.xlsx"
FIELD_REPORT = REPORT_DIR / "FIELD CHECK.xlsx"

# RAFIQ MD MOTIN UR — our head of branch, not an applicant.
SKIP_PASSPORTS = {"A00990004"}

PROGRAM_ORDER = ["KLP", "EAP", "BACHELOR", "MASTER"]
DEFAULT_BUDGET = 6          # students per pass


def _pid_alive(pid: int) -> bool:
    """Is that process still running?  A lock left behind by a process that died must not
    block the next run until it goes stale."""
    if pid <= 0:
        return False
    if os.name == "nt":
        import ctypes
        h = ctypes.windll.kernel32.OpenProcess(0x1000, False, pid)   # QUERY_LIMITED_INFORMATION
        if not h:
            return False
        code = ctypes.c_ulong()
        ctypes.windll.kernel32.GetExitCodeProcess(h, ctypes.byref(code))
        ctypes.windll.kernel32.CloseHandle(h)
        return code.value == 259                                     # STILL_ACTIVE
    try:
        os.kill(pid, 0)
        return True
    except OSError:
        return False


def _lock_held(path) -> bool:
    """True only when another *live* run holds this lock."""
    if not path.exists():
        return False
    if time.time() - path.stat().st_mtime >= LOCK_STALE_SECONDS:
        return False
    try:
        pid = int(path.read_text(encoding="utf-8", errors="replace").strip() or 0)
    except ValueError:
        return True                      # unreadable but fresh — leave it alone
    if _pid_alive(pid):
        return True
    logger.info("clearing a lock left by process %s, which is no longer running", pid)
    path.unlink(missing_ok=True)
    return False


# --- what has changed since last time ------------------------------------------------
def _hash(parts: List[str]) -> str:
    return hashlib.sha1("|".join(parts).encode("utf-8", "replace")).hexdigest()


def _file_parts(folder: Path) -> List[str]:
    out = []
    for f in sorted(folder.iterdir()):
        if f.is_file() and not f.name.startswith("."):
            st = f.stat()
            out.append(f"{f.name}:{st.st_size}:{int(st.st_mtime)}")
    return out


def doc_fingerprint(folder: Path) -> str:
    """Changes only when a document is added, replaced or resized — not when the portal is
    edited, so correcting a portal field never causes the scans to be read again."""
    return _hash(_file_parts(folder))


def field_fingerprint(folder: Path, student: Dict[str, str]) -> str:
    """Changes when a document changes OR the portal record is edited.

    The passport issue date lives in its own cache rather than the students export, so it
    has to be folded in here — otherwise correcting it on the portal would never cause the
    student to be checked again."""
    extra = ""
    try:
        from src.sheets import passport_issue
        extra = passport_issue.load().get((student.get("Passport No") or "").strip().upper(), "")
    except Exception:
        pass
    return _hash(_file_parts(folder) + [f"{k}={student[k]}" for k in sorted(student)]
                 + [f"_issue={extra}"])


def load_store() -> Dict[str, Dict[str, Any]]:
    if STORE_PATH.exists():
        try:
            data = json.loads(STORE_PATH.read_text(encoding="utf-8"))
            return {"documents": data.get("documents", {}), "fields": data.get("fields", {}),
                    "corrections": data.get("corrections", [])}
        except Exception as e:
            logger.warning("results store unreadable, starting fresh: %s", e)
    return {"documents": {}, "fields": {}, "corrections": []}


def save_store(store: Dict[str, Dict[str, Any]]) -> None:
    STORE_PATH.parent.mkdir(parents=True, exist_ok=True)
    tmp = STORE_PATH.with_suffix(".tmp")
    tmp.write_text(json.dumps(store, indent=1, ensure_ascii=False), encoding="utf-8")
    tmp.replace(STORE_PATH)


def pending(store: Dict[str, Dict[str, Any]], folders, portal) -> List[str]:
    """Passports that still need checking, oldest download first."""
    out = []
    for pas, (folder, _name) in folders.items():
        if pas in SKIP_PASSPORTS or pas not in portal:
            continue
        done_docs = store["documents"].get(pas, {}).get("fingerprint") == doc_fingerprint(folder)
        done_fields = store["fields"].get(pas, {}).get("fingerprint") == field_fingerprint(folder, portal[pas])
        if not (done_docs and done_fields):
            out.append((folder.stat().st_mtime, pas))
    return [p for _, p in sorted(out)]


# --- reading each document once -------------------------------------------------------
def _text_cache_path(pas: str) -> Path:
    return TEXT_DIR / f"{pas}.json"


def install_text_cache(pas: str) -> Dict[str, Any]:
    """Make both checkers read through a cache, so a document is only ever OCR'd once.

    The cache key is the file's name, size and modification time, so a re-uploaded or
    shrunk document is read again, and an unchanged one never is."""
    from src.verify import doc_verifier as dv
    from src.verify import field_check as fc

    path = _text_cache_path(pas)
    cache: Dict[str, Any] = {}
    if path.exists():
        try:
            cache = json.loads(path.read_text(encoding="utf-8"))
        except Exception as e:
            logger.warning("text cache for %s unreadable: %s", pas, e)

    original = dv._read_document_uncached if hasattr(dv, "_read_document_uncached") else dv.read_document
    dv._read_document_uncached = original

    def cached(file_path, max_pages: int = 6):
        f = Path(file_path)
        try:
            st = f.stat()
            key = f"{f.name}:{st.st_size}:{int(st.st_mtime)}:{max_pages}"
        except OSError:
            return original(file_path, max_pages)
        hit = cache.get(key)
        if hit is not None:
            return hit["text"], [tuple(x) for x in hit["sizes"]]
        text, sizes = original(file_path, max_pages)
        cache[key] = {"text": text, "sizes": [list(x) for x in sizes]}
        return text, sizes

    original_pages = (dv._read_pages_uncached if hasattr(dv, "_read_pages_uncached")
                      else dv.read_pages)
    dv._read_pages_uncached = original_pages

    def cached_pages(file_path, max_pages: int = 20, sideways=None):
        """Per-page text is as expensive as whole-file text and is now needed by three
        checks (NID, bank, academic), so it is cached the same way."""
        f = Path(file_path)
        try:
            st = f.stat()
            key = f"PAGES:{f.name}:{st.st_size}:{int(st.st_mtime)}:{max_pages}"
        except OSError:
            return original_pages(file_path, max_pages, sideways)
        hit = cache.get(key)
        if hit is not None:
            if sideways is not None:
                sideways.extend(hit.get("sideways", []))
            return list(hit["pages"])
        turned = []
        pages = original_pages(file_path, max_pages, turned)
        cache[key] = {"pages": pages, "sideways": list(turned)}
        if sideways is not None:
            sideways.extend(turned)
        return pages

    dv.read_document = cached
    fc.read_document = cached                 # field_check imported it by name
    dv.read_pages = cached_pages
    return {"cache": cache, "path": path, "original": original,
            "original_pages": original_pages}


def save_text_cache(handle: Dict[str, Any]) -> None:
    from src.verify import doc_verifier as dv
    from src.verify import field_check as fc
    try:
        TEXT_DIR.mkdir(parents=True, exist_ok=True)
        handle["path"].write_text(json.dumps(handle["cache"], ensure_ascii=False), encoding="utf-8")
    except Exception as e:
        logger.warning("could not save the text cache: %s", e)
    dv.read_document = handle["original"]
    fc.read_document = handle["original"]
    if handle.get("original_pages"):
        dv.read_pages = handle["original_pages"]


# --- checking one student ------------------------------------------------------------
def _record_corrections(store: Dict[str, Any], pas: str, program: str, name: str,
                        previous: List[Dict[str, Any]], current: List[Dict[str, Any]]) -> int:
    """Log every portal field that changed since this student was last checked, and say
    whether the change settled the problem."""
    was = {str(r.get("field")): r for r in previous}
    n = 0
    for r in current:
        old = was.get(str(r.get("field")))
        if not old:
            continue
        before, after = str(old.get("portal", "")), str(r.get("portal", ""))
        if before == after:
            continue
        store.setdefault("corrections", []).append({
            "noticed": dt.datetime.now().isoformat(timespec="seconds"),
            "student": name, "passport": pas, "program": program,
            "field": str(r.get("field")),
            "was": before, "now": after,
            "was_result": str(old.get("result", "")), "now_result": str(r.get("result", "")),
            "detail": str(r.get("detail", "")),
        })
        n += 1
    return n


def check_one(pas: str, folder: Path, student: Dict[str, str],
              store: Dict[str, Dict[str, Any]]) -> Tuple[str, str, int]:
    """Document check then, straight away, the progress-sheet check for the same student.
    Returns (verdict, field summary, number of portal corrections noticed)."""
    from src.verify import doc_verifier as dv
    from src.verify import field_check as fc

    program = dv.program_of(student)
    name = student.get("Full Name", "")
    doc_fp = doc_fingerprint(folder)
    field_fp = field_fingerprint(folder, student)

    handle = install_text_cache(pas)
    try:
        verdict = store["documents"].get(pas, {}).get("verdict", "")
        if store["documents"].get(pas, {}).get("fingerprint") != doc_fp:
            rows = dv.verify_student(folder, student, program)
            verdict = dv.student_verdict(rows)
            store["documents"][pas] = {
                "fingerprint": doc_fp, "checked": dt.datetime.now().isoformat(timespec="seconds"),
                "student": name, "program": program, "verdict": verdict, "rows": rows,
            }

        field_note, corrected = "", 0
        if store["fields"].get(pas, {}).get("fingerprint") != field_fp:
            previous = store["fields"].get(pas, {}).get("rows", [])
            frows = fc.check_student(pas, folder, student)
            if previous:
                corrected = _record_corrections(store, pas, program, name, previous, frows)
            differs = sum(1 for r in frows if r["result"] == fc.DIFFERS)
            match = sum(1 for r in frows if r["result"] == fc.MATCH)
            store["fields"][pas] = {
                "fingerprint": field_fp, "checked": dt.datetime.now().isoformat(timespec="seconds"),
                "student": name, "program": program, "rows": frows,
            }
            field_note = f"{match} matched, {differs} differ"
        return verdict, field_note, corrected
    finally:
        save_text_cache(handle)


def _cell(v):
    """Excel refuses control characters; nothing should ever fail a whole report."""
    if not isinstance(v, str):
        return v
    return "".join(c for c in v if ord(c) >= 32 or c in "\n\t")[:32000]


# --- the two reports -----------------------------------------------------------------
def _style_header(ws) -> None:
    from openpyxl.styles import Font
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"


def _programs_in(section: Dict[str, Any]) -> List[str]:
    found = {v.get("program", "") for v in section.values()}
    ordered = [p for p in PROGRAM_ORDER if p in found]
    return ordered + sorted(found - set(ordered) - {""})


def write_document_report(store: Dict[str, Dict[str, Any]], path: Path = DOC_REPORT) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill
    colours = {"PASS": "C6EFCE", "FLAG": "FFEB9C", "FAIL": "FFC7CE", "MISSING": "F2F2F2",
               "NOTE": "DDEBF7", "REVIEW": "FFEB9C", "INCOMPLETE": "F2F2F2"}
    section = store["documents"]
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    ws.append([_cell(x) for x in ["Student", "Passport", "Program", "Verdict", "Missing", "Fail", "Flag", "Pass", "Checked"]])
    _style_header(ws)
    for pas, rec in sorted(section.items(), key=lambda kv: (kv[1].get("program", ""), kv[1].get("student", ""))):
        counts: Dict[str, int] = {}
        for r in rec.get("rows", []):
            counts[r["verdict"]] = counts.get(r["verdict"], 0) + 1
        ws.append([rec.get("student", ""), pas, rec.get("program", ""), rec.get("verdict", ""),
                   counts.get("MISSING", 0), counts.get("FAIL", 0), counts.get("FLAG", 0),
                   counts.get("PASS", 0), rec.get("checked", "")[:16].replace("T", " ")])
        ws.cell(row=ws.max_row, column=4).fill = PatternFill(
            "solid", fgColor=colours.get(rec.get("verdict", ""), "FFFFFF"))
    for col, w in zip("ABCDEFGHI", (28, 13, 11, 13, 9, 7, 7, 7, 17)):
        ws.column_dimensions[col].width = w

    for program in _programs_in(section):
        sheet = wb.create_sheet(program[:31])
        sheet.append([_cell(x) for x in ["Student", "Passport", "Student verdict", "Document", "File", "Verdict", "Details"]])
        _style_header(sheet)
        for pas, rec in sorted(section.items(), key=lambda kv: kv[1].get("student", "")):
            if rec.get("program") != program:
                continue
            for r in rec.get("rows", []):
                sheet.append([_cell(x) for x in [rec.get("student", ""), pas, rec.get("verdict", ""),
                              r.get("doc", ""), r.get("file", ""), r.get("verdict", ""), r.get("detail", "")]])
                sheet.cell(row=sheet.max_row, column=6).fill = PatternFill(
                    "solid", fgColor=colours.get(r.get("verdict", ""), "FFFFFF"))
                sheet.cell(row=sheet.max_row, column=3).fill = PatternFill(
                    "solid", fgColor=colours.get(rec.get("verdict", ""), "FFFFFF"))
        for col, w in zip("ABCDEFG", (28, 13, 14, 38, 34, 10, 120)):
            sheet.column_dimensions[col].width = w
    wb.save(path)
    return path


def write_field_report(store: Dict[str, Dict[str, Any]], path: Path = FIELD_REPORT) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill
    from src.verify import field_check as fc
    colour = {fc.MATCH: "C6EFCE", fc.DIFFERS: "FFC7CE", fc.UNREADABLE: "FFEB9C",
              fc.NO_DOC: "F2F2F2", "BLANK": "FFFFFF"}
    section = store["fields"]
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    ws.append([_cell(x) for x in ["Student", "Passport", "Program", "Checked", "Match", "Differs",
               "Unreadable", "No document", "Blank", "Last checked"]])
    _style_header(ws)
    for pas, rec in sorted(section.items(), key=lambda kv: (kv[1].get("program", ""), kv[1].get("student", ""))):
        counts: Dict[str, int] = {}
        for r in rec.get("rows", []):
            counts[str(r["result"])] = counts.get(str(r["result"]), 0) + 1
        checked = counts.get(fc.MATCH, 0) + counts.get(fc.DIFFERS, 0) + counts.get(fc.UNREADABLE, 0)
        ws.append([rec.get("student", ""), pas, rec.get("program", ""), checked,
                   counts.get(fc.MATCH, 0), counts.get(fc.DIFFERS, 0), counts.get(fc.UNREADABLE, 0),
                   counts.get(fc.NO_DOC, 0), counts.get("BLANK", 0),
                   rec.get("checked", "")[:16].replace("T", " ")])
        if counts.get(fc.DIFFERS, 0):
            ws.cell(row=ws.max_row, column=6).fill = PatternFill("solid", fgColor="FFC7CE")
    for col, w in zip("ABCDEFGHIJ", (28, 13, 11, 9, 8, 8, 11, 13, 7, 17)):
        ws.column_dimensions[col].width = w

    for program in _programs_in(section):
        sheet = wb.create_sheet(program[:31])
        sheet.append([_cell(x) for x in ["Student", "Passport", "Field", "Portal value", "Result", "Detail"]])
        _style_header(sheet)
        for pas, rec in sorted(section.items(), key=lambda kv: kv[1].get("student", "")):
            if rec.get("program") != program:
                continue
            for r in rec.get("rows", []):
                sheet.append([_cell(x) for x in [rec.get("student", ""), pas, r.get("field", ""), r.get("portal", ""),
                              r.get("result", ""), r.get("detail", "")]])
                sheet.cell(row=sheet.max_row, column=5).fill = PatternFill(
                    "solid", fgColor=colour.get(str(r.get("result", "")), "FFFFFF"))
        for col, w in zip("ABCDEF", (28, 13, 22, 30, 12, 80)):
            sheet.column_dimensions[col].width = w

    # What changed on the portal after a check — usually an employee acting on a flag.
    ws3 = wb.create_sheet("Corrections")
    ws3.append([_cell(x) for x in ["Noticed", "Student", "Passport", "Program", "Field", "Was on portal",
                "Now on portal", "Was", "Now", "Detail"]])
    _style_header(ws3)
    for c in reversed(store.get("corrections", [])):       # newest first
        ws3.append([str(c.get("noticed", ""))[:16].replace("T", " "), c.get("student", ""),
                    c.get("passport", ""), c.get("program", ""), c.get("field", ""),
                    c.get("was", ""), c.get("now", ""), c.get("was_result", ""),
                    c.get("now_result", ""), c.get("detail", "")])
        ws3.cell(row=ws3.max_row, column=9).fill = PatternFill(
            "solid", fgColor=colour.get(str(c.get("now_result", "")), "FFFFFF"))
        ws3.cell(row=ws3.max_row, column=8).fill = PatternFill(
            "solid", fgColor=colour.get(str(c.get("was_result", "")), "FFFFFF"))
    for col, w in zip("ABCDEFGHIJ", (17, 28, 13, 11, 22, 26, 26, 12, 12, 70)):
        ws3.column_dimensions[col].width = w
    wb.save(path)
    return path


def write_reports(store: Dict[str, Dict[str, Any]]) -> List[Path]:
    return [write_document_report(store), write_field_report(store)]


# --- the pass the scheduler runs ------------------------------------------------------
def run(budget: int = DEFAULT_BUDGET, only: List[str] | None = None,
        recheck: bool = False) -> Dict[str, Any]:
    """Check what is pending (up to `budget` students) and rewrite both reports.

    `recheck` re-tests everybody against the current rules.  It happens inside the lock,
    because editing the stored results from outside would be overwritten by a pass that is
    already running.  Findings and the corrections history are kept; only the "already
    checked" marks are cleared, and the OCR text cache means nothing is read twice."""
    from src.verify import doc_verifier as dv
    from src.verify import field_check as fc

    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    if _lock_held(LOCK_PATH):
        print(f"Another check is already running (pid {LOCK_PATH.read_text(errors='replace').strip()}) — skipped.")
        return {"checked": [], "waiting": 0, "store": load_store()}
    LOCK_PATH.write_text(str(os.getpid()), encoding="utf-8")

    store = load_store()
    if recheck:
        for section in ("documents", "fields"):
            for rec in store[section].values():
                rec["fingerprint"] = ""
        print(f"re-checking all {len(store['documents'])} student(s) against the current rules")
    folders, portal = dv.student_folders(), dv.portal_students()
    queue = [p.upper() for p in only] if only else pending(store, folders, portal)
    waiting = len(queue)
    if budget:
        queue = queue[:budget]

    done: List[Dict[str, str]] = []
    skipped: List[str] = []
    try:
        for i, pas in enumerate(queue, 1):
            if pas not in folders or pas not in portal:
                continue
            student = portal[pas]
            name = student.get("Full Name", "")
            try:
                verdict, field_note, corrected = check_one(pas, folders[pas][0], student, store)
            except (NameError, AttributeError, ImportError, TypeError) as e:
                # A fault in the rules themselves affects every student.  Skipping them
                # one by one produces a report that looks clean and is worthless, so the
                # run stops and says so.  (A NameError in check_bank once did exactly
                # that: every bank check was skipped and nothing in the report showed it.)
                raise RuntimeError(
                    f"the rules failed on {name} ({pas}): {type(e).__name__}: {e}. "
                    "This is a code fault, not a document problem — the run has stopped "
                    "so the report is not silently wrong.") from e
            except Exception as e:
                logger.warning("%s (%s) could not be checked: %s", name, pas, e)
                skipped.append(f"{name} ({pas}): {e}")
                continue
            save_store(store)                  # a stopped run never loses finished work
            differs = sum(1 for r in store["fields"].get(pas, {}).get("rows", [])
                          if r["result"] == fc.DIFFERS)
            done.append({"student": name, "passport": pas, "program": dv.program_of(student),
                         "verdict": verdict, "differs": differs, "corrected": corrected})
            print(f"[{i}/{len(queue)}] {name[:30]:32} {verdict:11} {field_note}"
                  + (f" | {corrected} portal field(s) changed since last check" if corrected else ""),
                  flush=True)

        if done or not DOC_REPORT.exists() or not FIELD_REPORT.exists():
            write_reports(store)
    finally:
        LOCK_PATH.unlink(missing_ok=True)
    if skipped:
        print(f"\n{len(skipped)} student(s) could not be read:")
        for s in skipped[:10]:
            print(f"   {s}")
    return {"checked": done, "waiting": max(0, waiting - len(done)),
            "skipped": skipped, "store": store}


def _first_problem(result: Dict[str, Any], d: Dict[str, Any]) -> str:
    """Why a student FAILed (the first failing rule) or is INCOMPLETE (the first document not
    uploaded), from the stored document-check rows; "" when they are not at hand."""
    import re
    rows = (((result.get("store") or {}).get("documents") or {}).get(d.get("passport"), {}) or {}).get("rows") or []
    want = "MISSING" if d.get("verdict") == "INCOMPLETE" else "FAIL"
    for r in rows:
        if r.get("verdict") == want:
            detail = str(r.get("detail", ""))
            m = re.search(r"\bFAIL:\s*([^|]+)", detail) if want == "FAIL" else None
            detail = m.group(1).strip() if m else detail
            text = f"{r.get('doc', '')}: {detail}".strip(": ")
            return text if len(text) <= 160 else text[:159] + "…"
    return ""


def summary_lines(result: Dict[str, Any]) -> List[str]:
    """The verification part of the Telegram message: each student's verdict, with the first
    failing rule (or missing document) for a FAIL or INCOMPLETE."""
    done = result.get("checked", [])
    if not done:
        return []
    bad = [d for d in done if d["verdict"] in ("FAIL", "INCOMPLETE") or d["differs"]]
    out = [f"🔍 Documents checked: {len(done)}"]
    for d in done:
        why = _first_problem(result, d) if d["verdict"] in ("FAIL", "INCOMPLETE") else ""
        bits = [f"{d['verdict']} ({why})" if why else d["verdict"]]
        if d["differs"]:
            bits.append(f"{d['differs']} field(s) differ from the portal")
        if d.get("corrected"):
            bits.append(f"{d['corrected']} portal field(s) corrected since last check")
        out.append(f"   • {d['student']} ({d['program']}) — {', '.join(bits)}")
    if result.get("waiting"):
        out.append(f"   {result['waiting']} more waiting — they run on the next passes.")
    if bad:
        out.append(f"   Reports: DOCUMENT CHECK.xlsx / FIELD CHECK.xlsx in {REPORT_DIR}")
    return out


def main() -> None:
    import sys
    try:                                   # a plain console is cp1252 and chokes on emoji
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    logging.basicConfig(level=logging.WARNING, format="%(asctime)s  %(message)s")
    ap = argparse.ArgumentParser(description="Check downloaded documents and portal fields (local, no portal writes).")
    ap.add_argument("--budget", type=int, default=DEFAULT_BUDGET, help="students this pass (0 = no limit)")
    ap.add_argument("--passport", default="", help="check these passports only (comma separated)")
    ap.add_argument("--pending", action="store_true", help="list what is waiting and stop")
    ap.add_argument("--rebuild", action="store_true", help="only rewrite the two reports")
    ap.add_argument("--recheck", action="store_true",
                    help="re-test everyone against the current rules (keeps the corrections history)")
    ap.add_argument("--recheck-all", action="store_true", help="forget every stored result first")
    args = ap.parse_args()

    if args.recheck_all and STORE_PATH.exists():
        STORE_PATH.unlink()
        print("stored results cleared — everyone will be checked again")

    if args.rebuild:
        for p in write_reports(load_store()):
            print(f"report: {p}")
        return

    if args.pending:
        from src.verify import doc_verifier as dv
        store = load_store()
        q = pending(store, dv.student_folders(), dv.portal_students())
        print(f"{len(q)} student(s) waiting to be checked")
        for p in q[:20]:
            print(f"   {p}")
        return

    only = [p for p in args.passport.split(",") if p.strip()] or None
    result = run(budget=0 if only else args.budget, only=only, recheck=args.recheck)
    print()
    for line in summary_lines(result) or ["nothing was pending"]:
        print(line)
    print(f"\nreports: {DOC_REPORT}\n         {FIELD_REPORT}")


if __name__ == "__main__":
    main()
