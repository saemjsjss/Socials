"""The one-time copy of everything the bot has read to Supabase (D11), run by the owner after deploy.

    python -m src.cloud.backfill --dry-run       # write the payloads to data\\cloud\\dry_run\\ first
    python -m src.cloud.backfill                 # then for real

From the portal (read-only: GET plus the login POST, through portal_get, with a session of its
own): every page of students.php (student, and verification from the same list), the CSV export
(student_export), every student's progress.php, 4 at a time as /stage does (student_progress),
the verified-documents list (student_documents), consult_requests.php for every day from the
oldest request the portal's date filter finds up to today (consultation, consultation_day) and
its all-time totals, every page of students.php?status=pending (pending_payment), the window
applications under review, index.php (dashboard_fact) and calendar.php (calendar_item). From
disk: data\\verification\\results.json (doc_verdict, doc_check, field_check, field_correction and
the DOCUMENT CHECK / FIELD CHECK reports), every OCR text cache data\\verification\\text\\*.json
(doc_page_text), the passport watcher's memory (passport_alert), data\\passport_issue.json
(passport_issue) and the newest missing-information report (report, report_section).

It prints its progress per kind (counts only, never student data). A read that fails is said
and skipped; a partial read deletes nothing. Refuses to start in the quiet windows of 01 §5.4
(18:00-18:10, 08:25-08:40, 09:00-09:10) or while a portal sync runs, unless --force.

The collect_* functions return records.batch()es and the failed reads' reasons, so the hourly
"full picture" job reads the same way (collect_students, collect_pending, collect_consultations
for today and yesterday, collect_totals, collect_window_applications, collect_dashboard,
collect_calendar), then hands them to src.cloud.handoff.submit("full_picture", ...).
"""
from __future__ import annotations

import os
import sys

if __name__ == "__main__":          # a CPU-only embedding process: before torch can be imported
    os.environ["CUDA_VISIBLE_DEVICES"] = ""

import argparse
import asyncio
import csv
import io
import json
import logging
import re
import time
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Sequence, Tuple

from src.cloud import records
from src.config import BOT_ROOT, settings

logger = logging.getLogger("hangeul.cloud")

Batches = List[Dict[str, Any]]
CONSULT_FLOOR = date(2015, 1, 1)       # the oldest day the search for the first request starts from
# Quiet windows (Asia/Dhaka) of 01 §5.4: scheduled jobs run then.
QUIET_WINDOWS = ((18, 0, 18, 10), (8, 25, 8, 40), (9, 0, 9, 10))
SYNC_LOCK_STALE = 2 * 3600


def _why(e: Exception) -> str:
    from src.scraper.client import portal_error_reason
    return portal_error_reason(e)


def quiet_reason(at: Optional[datetime] = None, data_dir: Optional[Path] = None) -> Optional[str]:
    """Why now is no time to read the portal in bulk (None when it is): a quiet window, or a
    portal sync running (data/auto_sync.lock held by a live process; the lock is only read)."""
    at = at or records.now()
    minutes = at.hour * 60 + at.minute
    for h1, m1, h2, m2 in QUIET_WINDOWS:
        if h1 * 60 + m1 <= minutes < h2 * 60 + m2:
            return f"{h1:02d}:{m1:02d}-{h2:02d}:{m2:02d} is a quiet window (a scheduled job runs then)"
    lock = (data_dir or BOT_ROOT / "data") / "auto_sync.lock"
    try:
        if lock.exists() and time.time() - lock.stat().st_mtime < SYNC_LOCK_STALE:
            from src.sheets.auto_sync import _pid_alive
            pid = int(lock.read_text(encoding="utf-8", errors="replace").strip() or 0)
            if _pid_alive(pid):
                return "a portal sync is running (data/auto_sync.lock)"
    except (OSError, ValueError):
        pass
    return None


# --------------------------------------------------------------------------- portal collectors

async def collect_students(client, today: date) -> Tuple[Batches, List[str], Optional[List[Dict[str, Any]]]]:
    """Every page of students.php -> the student batch (complete) and the verification batches
    (per stamp day; complete within the last year when every stamp could be read), the failed
    reads, and the students themselves (None when the list could not be read)."""
    try:
        students = await client.read_students()
    except Exception as e:
        return [], [f"students.php: {_why(e)}"], None
    at = records.as_read_at(None)
    out = [records.batch("student", "all", records.students(students, at), True)]
    lines = [s for s in students if s.get("verified_line")]
    from src.dates import parse_stamp
    stamps_ok = (not students or bool(lines)) and all(parse_stamp(s.get("verified_stamp") or "") for s in lines)
    out.append(records.batches("verification", records.verifications(students, today, at), stamps_ok,
                               records.verification_window(today) if stamps_ok else None))
    failed = [] if stamps_ok else ["students.php: a verification stamp could not be read (layout not recognised)"]
    return out, failed, students


async def collect_export(client, total: Optional[int] = None) -> Tuple[Batches, List[str]]:
    """students.php?export=csv (every column, B2B and Direct) -> the student_export batch. Complete
    when the export came whole (its header has Student ID and Full Name) and holds at least 90 % of
    the list's own total (`total`, when known): the export has no pager to prove it by."""
    try:
        resp = await client.portal_get("students.php", params={"export": "csv"}, timeout=60.0)
        text = resp.content.decode("utf-8-sig", errors="replace")
        if "text/csv" not in resp.headers.get("content-type", "") and "Full Name" not in text[:2000]:
            return [], ["students.php?export=csv: the portal did not send the export"]
        rows = [dict(r) for r in csv.DictReader(io.StringIO(text))]
    except Exception as e:
        return [], [f"students.php?export=csv: {_why(e)}"]
    header_ok = bool(rows) and {"Student ID", "Full Name"} <= set(rows[0])
    complete = header_ok and (total is None or len(rows) >= 0.9 * total)
    return [records.batch("student_export", "all", records.student_exports(rows), complete)], \
        ([] if complete else ["students.php?export=csv: fewer rows than the student list, or no header"])


async def collect_documents(client) -> Tuple[Batches, List[str]]:
    """students.php?source=direct&filter_docs=verified, every page -> student_documents (complete)."""
    from src.sheets.verified_docs import fetch_verified_students
    try:
        rows = await fetch_verified_students(client)
    except Exception as e:
        return [], [f"students.php?source=direct&filter_docs=verified: {_why(e)}"]
    return [records.batch("student_documents", "all", records.student_documents(rows), True)], []


async def collect_progress(students: Sequence[Dict[str, Any]]) -> Tuple[Batches, List[str]]:
    """Every student's progress.php, 4 at a time with a session of its own (stage_report.read_progress,
    in a worker thread) -> student_progress, complete only when every page was read."""
    from src.sheets.stage_report import read_progress
    uids = [str(s["uid"]) for s in students if str(s.get("uid") or "").isdigit()]
    if not uids:
        return [], []
    try:
        pages = await asyncio.to_thread(read_progress, uids)
    except Exception as e:
        return [], [f"progress.php: {_why(e)}"]
    errors = [p.get("error") for p in pages.values() if isinstance(p, dict) and p.get("error")]
    listed = {str(s.get("uid")): s for s in students}
    complete = not errors and len(pages) == len(uids)
    rows = records.student_progress(pages, listed, records.as_read_at(None))
    failed = [f"progress.php: {len(errors)} of {len(uids)} pages not read ({errors[0]})"] if errors else []
    return [records.batch("student_progress", "all", rows, complete)], failed


async def _range_count(client, first: date, last: date) -> int:
    """How many requests the portal counts from `first` to `last` (its date filter's All tab)."""
    from src.scraper.client import PortalUnavailable
    view = await client.read_consultation_view({"status": "all", "from": first.isoformat(), "to": last.isoformat()})
    if view["from"] != first.isoformat() or view["to"] != last.isoformat():
        raise PortalUnavailable("consult_requests.php did not apply the date filter (layout not recognised)")
    return int(view["tabs"]["All"])


async def oldest_consultation_day(client, today: date, floor: date = CONSULT_FLOOR) -> Optional[date]:
    """The day of the oldest request the portal's date filter finds (a search on its own counts,
    ~13 reads), or None when it counts none since `floor`."""
    if await _range_count(client, floor, today) == 0:
        return None
    lo, hi = floor, today
    while lo < hi:
        mid = lo + timedelta(days=(hi - lo).days // 2)
        if await _range_count(client, floor, mid) > 0:
            hi = mid
        else:
            lo = mid + timedelta(days=1)
    return lo


async def collect_consultations(client, days: Sequence[date], *,
                                whole_range: bool = False) -> Tuple[Batches, List[str]]:
    """consult_requests.php's date filter for each day -> one consultation batch per day (complete
    when the portal listed every request of the day) and the consultation_day batch (the days'
    counts; complete only with `whole_range`, a read of every day there is, as the backfill does:
    a read of some days must not delete the others). Stops at the first unreachable answer."""
    out: Batches = []
    failed: List[str] = []
    day_rows: List[Dict[str, Any]] = []
    for day in days:
        try:
            on_day = await client.read_consultation_day(day)
        except Exception as e:
            failed.append(f"consult_requests.php {day.isoformat()}: {_why(e)}")
            if getattr(e, "unreachable", False):
                break
            continue
        at = records.as_read_at(None)
        rows, complete = records.consultations(on_day, at)
        out.append(records.batch("consultation", day.isoformat(), rows, complete))
        rec = records.consultation_day(on_day, at)
        if rec is not None:
            day_rows.append(rec)
    if day_rows:
        out.append(records.batch("consultation_day", "all", day_rows, whole_range and not failed))
    return out, failed


async def collect_totals(client) -> Tuple[Batches, List[str]]:
    try:
        totals = await client.read_consultation_totals()
    except Exception as e:
        return [], [f"consult_requests.php?status=file_opened: {_why(e)}"]
    rec = records.consultation_totals(totals)
    return ([records.batch("consultation_totals", "all", [rec], True)] if rec else []), []


async def collect_pending(client) -> Tuple[Batches, List[str]]:
    """Every page of students.php?status=pending -> pending_payment (rows whose own Payment is
    Pending, with the badge), complete."""
    from src.scraper.parsers import StudentListLayoutError, parse_pending_payments, parse_students_page
    try:
        pages = await client.read_student_pages({"status": "pending"})
        badge = (await asyncio.to_thread(parse_pending_payments, pages[0]) or {}).get("badge") if pages else None
        rows = []
        for html in pages:
            try:
                rows += (await asyncio.to_thread(parse_students_page, html))["students"]
            except StudentListLayoutError as e:
                return [], [f"students.php?status=pending: {e}"]
    except Exception as e:
        return [], [f"students.php?status=pending: {_why(e)}"]
    return [records.batch("pending_payment", "all", records.pending_payments(rows, badge), True)], []


_PAGED_RE = re.compile(r"Page\s+\d+\s+of\s+\d+|[?&](?:pg|page)=\d", re.I)


async def collect_window_applications(client) -> Tuple[Batches, List[str]]:
    """window_applications.php?status=under_review -> window_application; complete when the page
    has no pager (it has no total to prove more pages by)."""
    from src.scraper.parsers import parse_window_applications
    try:
        html = await client.fetch_html("window_applications.php?status=under_review")
    except Exception as e:
        return [], [f"window_applications.php: {_why(e)}"]
    rows = await asyncio.to_thread(parse_window_applications, html)
    if rows is None:
        return [], ["window_applications.php: no table with a Status column (layout not recognised)"]
    return [records.batch("window_application", "all", records.window_applications(rows),
                          not _PAGED_RE.search(html))], []


async def collect_dashboard(client) -> Tuple[Batches, List[str]]:
    """index.php -> dashboard_fact (ask.dashboard_facts: tiles and cards); complete when every tile
    group and card is on the page."""
    from src.bot.ask import dashboard_facts
    try:
        html = await client.fetch_html("index.php", timeout=30.0)
    except Exception as e:
        return [], [f"index.php: {_why(e)}"]
    facts = await asyncio.to_thread(dashboard_facts, html)
    if not facts:
        return [], ["index.php: no figure could be read (layout not recognised)"]
    return [records.batch("dashboard_fact", "all", records.dashboard_facts(facts),
                          records.dashboard_complete(facts))], []


async def collect_calendar(client, today: date) -> Tuple[Batches, List[str]]:
    """calendar.php -> calendar_item (ask.calendar_items), never complete (the page shows only this
    month and the next 45 days)."""
    from src.bot.ask import calendar_items
    try:
        html = await client.fetch_html("calendar.php")
    except Exception as e:
        return [], [f"calendar.php: {_why(e)}"]
    items, ok = await asyncio.to_thread(calendar_items, html, today)
    if not ok:
        return [], ["calendar.php: layout not recognised"]
    return [records.batch("calendar_item", "all", records.calendar_items(items), False)], []


# --------------------------------------------------------------------------- disk collectors

def _read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def _mtime(path: Path) -> str:
    return records.as_read_at(datetime.fromtimestamp(path.stat().st_mtime))


ALL_SCOPES = ("", chr(0xFFFF))      # every scope: the store holds every passport there is


def collect_results(verification_dir: Path) -> Tuple[Batches, List[str], Dict[str, Any]]:
    """results.json -> doc_verdict and field_check (per passport, complete), doc_check (complete),
    field_correction (append-only) and the DOCUMENT CHECK / FIELD CHECK reports; and the store."""
    path = verification_dir / "results.json"
    try:
        store = _read_json(path)
        if not isinstance(store, dict):
            raise ValueError("not an object")
    except FileNotFoundError:
        return [], [], {}
    except Exception as e:
        return [], [f"results.json: unreadable ({type(e).__name__})"], {}
    docs = {p: e for p, e in (store.get("documents") or {}).items() if isinstance(e, dict)}
    fields = {p: e for p, e in (store.get("fields") or {}).items() if isinstance(e, dict)}
    verdicts = [r for p, e in docs.items() for r in records.doc_verdicts(p, e)]
    checks = [r for p, e in fields.items() for r in records.field_checks(p, e)]
    out = [records.batches("doc_verdict", verdicts, True, ALL_SCOPES),
           records.batches("field_check", checks, True, ALL_SCOPES),
           records.batch("doc_check", "all", records.doc_checks(docs, fields), True),
           records.batch("field_correction", "all", records.field_corrections(store.get("corrections") or []), False)]
    for maker in (records.document_check_report, records.field_check_report):
        whole, sections = maker(store)
        if whole is not None:
            out.append(records.batch("report", whole["scope"], [whole], False))
            out.append(records.batch("report_section", f"{whole['key']}", sections, True))
    return out, [], store


def _current_files(docs_root: Path, passport: str) -> Dict[str, Tuple[int, int]]:
    """{file name: (size, mtime)} of the files now in the student's document folder
    (<DOCS_ROOT>\\<PROGRAM>\\<NAME (PASSPORT)>\\), {} when there is none."""
    out: Dict[str, Tuple[int, int]] = {}
    if not docs_root.exists():
        return out
    for folder in docs_root.glob(f"*/*({passport})"):
        if folder.is_dir():
            for f in folder.iterdir():
                if f.is_file() and not f.name.startswith("."):
                    st = f.stat()
                    out.setdefault(f.name, (st.st_size, int(st.st_mtime)))
    return out


def collect_page_texts(verification_dir: Path, store: Dict[str, Any], docs_root: Path) -> Tuple[Batches, List[str]]:
    """Every data/verification/text/<PASSPORT>.json -> doc_page_text, one batch per passport
    (complete: the cache file is the whole of it)."""
    names = {p: (e or {}).get("student", "") for p, e in (store.get("documents") or {}).items()}
    out: Batches = []
    failed: List[str] = []
    text_dir = verification_dir / "text"
    for path in sorted(text_dir.glob("*.json")) if text_dir.exists() else []:
        passport = path.stem
        try:
            cache = _read_json(path)
        except Exception as e:
            failed.append(f"an OCR text cache file: unreadable ({type(e).__name__})")
            continue
        rows = records.doc_page_texts(passport, cache, names.get(passport, ""),
                                      _current_files(docs_root, passport), _mtime(path))
        scope = records._passport(passport) or passport
        out.append(records.batch("doc_page_text", scope, rows, True))
    return out, failed


def collect_watcher(data_dir: Path) -> Tuple[Batches, List[str]]:
    """data/alerted_passport_issues.json (v2) -> passport_alert, complete (the memory is the truth)."""
    path = data_dir / "alerted_passport_issues.json"
    try:
        memory = _read_json(path)
    except FileNotFoundError:
        return [], []
    except Exception as e:
        return [], [f"alerted_passport_issues.json: unreadable ({type(e).__name__})"]
    if not isinstance(memory, dict) or memory.get("version") != 2 or not isinstance(memory.get("scans"), dict):
        return [], ["alerted_passport_issues.json: not the version-2 memory"]
    return [records.batch("passport_alert", "all", records.passport_alerts(memory), True)], []


def collect_issue_dates(data_dir: Path) -> Tuple[Batches, List[str]]:
    """data/passport_issue.json -> passport_issue, complete (the whole file)."""
    path = data_dir / "passport_issue.json"
    try:
        by_passport = _read_json(path).get("by_passport")
        if not isinstance(by_passport, dict):
            raise ValueError("no by_passport")
    except FileNotFoundError:
        return [], []
    except Exception as e:
        return [], [f"passport_issue.json: unreadable ({type(e).__name__})"]
    return [records.batch("passport_issue", "all", records.passport_issues(by_passport, _mtime(path)), True)], []


def collect_missing_report(data_dir: Path) -> Tuple[Batches, List[str]]:
    """The newest data/missing_reports/missing_information_YYYY-MM-DD.xlsx (its rows: the data the
    report is built from; the file itself is never uploaded) -> report + report_section."""
    files = sorted((data_dir / "missing_reports").glob("missing_information_*.xlsx"))
    if not files:
        return [], []
    path = files[-1]
    m = re.search(r"(\d{4}-\d{2}-\d{2})", path.name)
    if not m:
        return [], []
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = [tuple(r) for r in wb.worksheets[0].iter_rows(min_row=2, values_only=True) if any(r)]
        finally:
            wb.close()
    except Exception as e:
        return [], [f"{path.name}: unreadable ({type(e).__name__})"]
    whole, sections = records.missing_report(m.group(1), rows, read_at=_mtime(path), source=path.name)
    return [records.batch("report", "missing_report", [whole], False),
            records.batch("report_section", whole["key"], sections, True)], []


# --------------------------------------------------------------------------- the run

def _say(results, kind: str, complete: bool, dry: bool, say: Callable[[str], None]) -> None:
    rows = sum(r.rows for r in results)
    changed = sum(r.changed for r in results)
    calls = sum(r.calls for r in results)
    failed = [r.error for r in results if not r.ok]
    scopes = len(results)
    where = f" in {scopes} scopes" if scopes > 1 else ""
    if dry:
        line = f"{rows} record(s){where}, {changed} to send, {calls} request(s) written"
    else:
        line = (f"{rows} record(s){where}, {changed} sent: {sum(r.upserted for r in results)} upserted, "
                f"{sum(r.unchanged for r in results)} unchanged, {sum(r.deleted for r in results)} deleted")
    if not complete:
        line += " (not a complete read: nothing deleted)"
    if failed:
        line += f"; FAILED: {failed[0]}"
    say(f"  {kind}: {line}")


def publish_all(batches: Batches, say: Callable[[str], None] = print, only: Optional[set] = None) -> List[Any]:
    """Publish each batch now (in the current run), saying its counts."""
    from src.cloud import publish
    results = []
    for b in batches:
        if not b or (only and b["kind"] not in only):
            continue
        if b.get("scope") is None:
            res = publish.publish_scopes(b["kind"], b["rows"], b["complete"], b.get("scope_range"))
        else:
            res = [publish.publish(b["kind"], b["scope"], b["rows"], b["complete"], all_keys=b.get("all_keys"))]
        _say(res, b["kind"], b["complete"], publish._dry(None), say)
        results += res
    return results


async def _portal(args, today: date, say: Callable[[str], None], failed: List[str], only: Optional[set]) -> None:
    from src.scraper.client import HangeulAdminClient
    client = HangeulAdminClient()
    try:
        def step(what: str) -> None:
            say(f"Reading {what} ...")

        step("every page of students.php")
        b, f, students = await collect_students(client, today)
        failed += f
        say(f"  {len(students)} students" if students is not None else f"  not read: {f[0]}")
        publish_all(b, say, only)
        if students is not None and (not only or "student_progress" in only):
            step(f"every student's progress.php ({len(students)} pages, 4 at a time)")
            b, f = await collect_progress(students)
            failed += f
            publish_all(b, say, only)

        step("students.php?export=csv")
        b, f = await collect_export(client, len(students) if students else None)
        failed += f
        publish_all(b, say, only)

        step("the verified-documents list")
        b, f = await collect_documents(client)
        failed += f
        publish_all(b, say, only)

        if not only or {"consultation", "consultation_day"} & only:
            if args.days:
                days = [today - timedelta(days=n) for n in range(args.days - 1, -1, -1)]
                whole = False
            else:
                step("consult_requests.php for its oldest request")
                try:
                    first = await oldest_consultation_day(client, today)
                except Exception as e:
                    failed.append(f"consult_requests.php: {_why(e)}")
                    first = None
                    say(f"  not read: {_why(e)}")
                days = [first + timedelta(days=n) for n in range((today - first).days + 1)] if first else []
                whole = first is not None
                if first:
                    say(f"  the oldest request is from {first:%d %b %Y}: {len(days)} days to read")
            if days:
                step(f"consult_requests.php day by day ({len(days)} days)")
                b, f = await collect_consultations(client, days, whole_range=whole)
                failed += f
                if f:
                    say(f"  {len(f)} day(s) not read, e.g. {f[0]}")
                cons = [x for x in b if x["kind"] == "consultation"]
                if cons and (not only or "consultation" in only):
                    from src.cloud import publish
                    res = [publish.publish(x["kind"], x["scope"], x["rows"], x["complete"]) for x in cons]
                    _say(res, "consultation", all(x["complete"] for x in cons), publish._dry(None), say)
                publish_all([x for x in b if x["kind"] != "consultation"], say, only)

        for what, collect in (("the consultation totals", lambda: collect_totals(client)),
                              ("every page of students.php?status=pending", lambda: collect_pending(client)),
                              ("window_applications.php?status=under_review",
                               lambda: collect_window_applications(client)),
                              ("index.php", lambda: collect_dashboard(client)),
                              ("calendar.php", lambda: collect_calendar(client, today))):
            step(what)
            b, f = await collect()
            failed += f
            if f:
                say(f"  not read: {f[0]}")
            publish_all(b, say, only)
    finally:
        await client.close()


def _disk(args, say: Callable[[str], None], failed: List[str], only: Optional[set]) -> None:
    data_dir = Path(args.data_dir) if args.data_dir else BOT_ROOT / "data"
    verification_dir = Path(args.verification_dir) if args.verification_dir else settings.verification_dir()
    docs_root = Path(args.docs_root) if args.docs_root else settings.docs_root()
    say(f"Reading {verification_dir / 'results.json'} ...")
    b, f, store = collect_results(verification_dir)
    failed += f
    publish_all(b, say, only)
    if not only or "doc_page_text" in only:
        say(f"Reading the OCR text caches in {verification_dir / 'text'} ...")
        b, f = collect_page_texts(verification_dir, store, docs_root)
        failed += f
        from src.cloud import publish
        res = [publish.publish(x["kind"], x["scope"], x["rows"], x["complete"]) for x in b]
        _say(res, "doc_page_text", True, publish._dry(None), say)
    for what, collect in (("the passport watcher's memory", lambda: collect_watcher(data_dir)),
                          ("data/passport_issue.json", lambda: collect_issue_dates(data_dir)),
                          ("the newest missing-information report", lambda: collect_missing_report(data_dir))):
        say(f"Reading {what} ...")
        b, f = collect()
        failed += f
        if f:
            say(f"  not read: {f[0]}")
        publish_all(b, say, only)


def main(argv: Optional[Sequence[str]] = None) -> int:
    ap = argparse.ArgumentParser(description="Copy everything the bot reads to Supabase, once (read-only portal).")
    ap.add_argument("--dry-run", action="store_true", help="write the payloads to data/cloud/dry_run/, send nothing")
    ap.add_argument("--only", default="", help="comma-separated kinds (default: every kind)")
    ap.add_argument("--skip-portal", action="store_true", help="publish only what is on disk")
    ap.add_argument("--skip-disk", action="store_true", help="publish only what the portal shows")
    ap.add_argument("--days", type=int, default=0,
                    help="consultations: only the last N days (default: every day from the oldest request)")
    ap.add_argument("--ignore-state", action="store_true",
                    help="send every record, not only those that changed (Supabase answers 'unchanged' for the same)")
    ap.add_argument("--data-dir", default="", help="read the watcher memory, issue dates and missing reports here")
    ap.add_argument("--verification-dir", default="", help="read results.json and text/ here")
    ap.add_argument("--docs-root", default="", help="the student document folders (to match OCR cache versions)")
    ap.add_argument("--force", action="store_true", help="run even in a quiet window or during a portal sync")
    args = ap.parse_args(argv)

    from src.cloud import embed, publish
    say = print
    if not embed.custom_embedder():
        try:
            embed.prepare_process()          # CPU only, offline; a no-op when already done
        except embed.EmbedError as e:
            say(f"Cannot embed in this process: {e}. Run it as python -m src.cloud.backfill.")
            return 2
    if not args.dry_run and not publish.enabled():
        say("Publishing is off: set SUPABASE_URL, SUPABASE_SECRET_KEY and CLOUD_PUBLISH_ENABLED=true in .env "
            "(or use --dry-run). Nothing was read.")
        return 2
    reason = quiet_reason(data_dir=Path(args.data_dir) if args.data_dir else None)
    if reason and not args.force and not args.skip_portal:
        say(f"Not now: {reason}. Try again in a few minutes (or --force).")
        return 2
    only = {k.strip() for k in args.only.split(",") if k.strip()} or None
    publish.set_dry_run(args.dry_run)
    if args.ignore_state and not args.dry_run and publish.STATE_PATH.exists():
        backup = publish.STATE_PATH.with_name(publish.STATE_PATH.name + ".bak")
        os.replace(publish.STATE_PATH, backup)
        say(f"The hash state was moved to {backup.name}: every record is sent.")
    started = time.perf_counter()
    today = records.now().date()
    failed: List[str] = []
    with publish.publisher_lock() as got:
        if not got:
            say("Another publisher is running; try again later.")
            return 1
        with publish.run("backfill", dry_run=args.dry_run) as r:
            if not args.skip_portal:
                asyncio.run(_portal(args, today, say, failed, only))
            if not args.skip_disk:
                _disk(args, say, failed, only)
            if r is not None:
                r.failed_reads.extend(failed)
    say(f"Done in {time.perf_counter() - started:.0f} s; {len(failed)} read(s) failed."
        + (f" Payloads are in {r.dry_dir}." if args.dry_run and r is not None and r.dry_dir else ""))
    for f in failed[:20]:
        say(f"  not read: {f}")
    return 0


if __name__ == "__main__":
    from src.cloud import embed as _embed
    _embed.prepare_process()
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    logging.basicConfig(level=logging.WARNING, format="%(asctime)s [%(levelname)s] %(name)s: %(message)s")
    from src.cloud import backfill as _backfill
    sys.exit(_backfill.main())
